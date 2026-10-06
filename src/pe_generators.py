from __future__ import annotations

from datetime import date, timedelta
from io import BytesIO
from pathlib import Path

import numpy as np
import pandas as pd
import yaml
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.cell import WriteOnlyCell

from . import generators as g


GL_COLUMNS = [
    "S.No", "Reference Number", "Batch Number", "Vendor", "Affiliate",
    "Deal Name", "Position", "Bank Account", "Investor Name", "Amount",
    "Debit", "Credit", "Location ID", "Fund Name",
]

SOI_COLUMNS = [
    "S.No", "Fund Name", "Fund ID", "Deal Name", "Investment ID", "Position",
    "Security Type", "Investor Name", "Investment Date", "Acquisition Date",
    "Currency", "Beginning Cost", "Additions", "Disposals", "Ending Cost",
    "Beginning Fair Value", "Unrealized Gain/Loss", "Realized Gain/Loss",
    "FX Gain/Loss", "Ending Fair Value", "Ownership %", "Shares/Units",
    "Valuation Method", "Investment Status", "Reporting Date",
]

ROLLFORWARD_COLUMNS = [
    "S.No", "Fund Name", "Fund ID", "Deal Name", "Investment ID", "Position",
    "Investor Name", "Reporting Period", "Beginning Cost", "Purchases/Additions",
    "Follow-on Investment", "PIK/Capitalized Interest", "Disposals", "Write-offs",
    "Transfers/Adjustments", "Other Adjustments", "Ending Cost", "Beginning Fair Value",
    "Interest Income", "Beginning Accrued Interest", "Ending Accrued Interest", "Unrealized Gain/Loss",
    "Realized Gain/Loss", "FX Gain/Loss", "Write-up", "Write-down",
    "Ending Fair Value", "Source Transaction Reference",
]

TRANSACTION_COLUMNS = [
    "Reference Number", "Batch Number", "Transaction Type", "Fund Name", "Fund ID",
    "Deal Name", "Investment ID", "Position", "Investor Name", "Reporting Period",
    "Transaction Amount", "Cost Change", "Fair Value Change", "Realized Gain/Loss",
    "Carrying Value", "Affiliate", "Vendor", "Bank Account", "Location ID",
]


def _master_names(config_dir: Path, module: str, candidates: tuple[str, ...], count: int, seed: int) -> np.ndarray:
    path = config_dir / f"{module}.yaml"
    if path.exists():
        with path.open("r", encoding="utf-8") as handle:
            schema = yaml.safe_load(handle) or {}
        by_name = {column.get("name", "").lower(): column for column in schema.get("columns", [])}
        column = next((by_name[name] for name in candidates if name in by_name), None)
        if column is not None and column.get("type") == "name_pool":
            params = {key: value for key, value in column.items() if key not in {"name", "type"}}
            return g.gen_name_pool(count, g.get_rng(seed), **params)
    return g.gen_name_pool(count, g.get_rng(seed), pool="company", pool_size=25000, unique=True)


def _period_dates(reporting_date: date, count: int) -> list[date]:
    return [reporting_date - timedelta(days=92 * (count - index - 1)) for index in range(count)]


def _numbered(prefix: str, start: int, count: int, width: int = 10) -> np.ndarray:
    numbers = np.arange(start, start + count, dtype=np.int64)
    return np.char.mod(f"{prefix}%0{width}d", numbers).astype(object)


def _base_transactions(
    refs: np.ndarray,
    batches: np.ndarray,
    kind: np.ndarray,
    fund_name: np.ndarray,
    fund_id: np.ndarray,
    deal_name: np.ndarray,
    investment_id: np.ndarray,
    position: np.ndarray,
    investor: np.ndarray,
    period: np.ndarray,
    amount: np.ndarray,
    cost_change: np.ndarray,
    fv_change: np.ndarray,
    realized: np.ndarray,
    carrying: np.ndarray,
    affiliate: np.ndarray,
    vendor: np.ndarray,
    bank: np.ndarray,
    location: np.ndarray,
) -> pd.DataFrame:
    return pd.DataFrame({
        "Reference Number": refs, "Batch Number": batches, "Transaction Type": kind,
        "Fund Name": fund_name, "Fund ID": fund_id, "Deal Name": deal_name,
        "Investment ID": investment_id, "Position": position, "Investor Name": investor,
        "Reporting Period": period, "Transaction Amount": amount, "Cost Change": cost_change,
        "Fair Value Change": fv_change, "Realized Gain/Loss": realized,
        "Carrying Value": carrying, "Affiliate": affiliate, "Vendor": vendor,
        "Bank Account": bank, "Location ID": location,
    }, columns=TRANSACTION_COLUMNS)


def generate_linked_datasets(
    record_count: int,
    fund_count: int,
    investment_count: int,
    period_count: int,
    reporting_date: date,
    seed: int = 42,
    base_currency: str = "USD",
    config_dir: Path | None = None,
    profile: str = "PRIVATE_EQUITY",
) -> dict[str, pd.DataFrame | dict]:
    """Generate one transaction model and derive its SOI, rollforward, and GL."""
    if record_count < 1 or fund_count < 1 or investment_count < 1 or period_count < 1:
        raise ValueError("Record, fund, investment, and reporting-period counts must be at least one.")
    if record_count > 1_000_000:
        raise ValueError("SOI record count cannot exceed 1,000,000.")
    if period_count > 20:
        raise ValueError("Reporting periods cannot exceed 20.")
    if investment_count > record_count:
        raise ValueError("Number of investments cannot exceed the SOI record count.")
    if fund_count > investment_count:
        raise ValueError("Number of funds cannot exceed the number of investments.")

    rng = g.get_rng(seed)
    config_dir = config_dir or Path(__file__).resolve().parent.parent / "config"
    deal_count = investment_count
    ref_investors = _master_names(config_dir, "investors", ("investor_name", "investor name"), min(record_count, 5000), seed + 1)
    ref_vendors = _master_names(config_dir, "vendors", ("vendor_name", "vendor name"), 500, seed + 2)
    ref_affiliates = _master_names(config_dir, "affiliates", ("affiliate_name", "affiliate name"), max(fund_count, 1), seed + 3)
    fund_names = g.gen_name_pool(fund_count, rng, pool="company", pool_size=25000, unique=True)
    deal_names = g.gen_name_pool(deal_count, rng, pool="company", pool_size=25000, unique=True)
    fund_ids = _numbered("FUND-", 1, fund_count, 5)
    deal_funds = np.arange(deal_count) % fund_count
    investment_ids = _numbered("INV-", 1, record_count, 9)
    deal_indexes = np.arange(record_count) % deal_count
    fund_indexes = deal_funds[deal_indexes]
    investor_indexes = rng.integers(0, len(ref_investors), size=record_count)
    affiliate_indexes = rng.integers(0, len(ref_affiliates), size=record_count)
    vendor_indexes = rng.integers(0, len(ref_vendors), size=record_count)
    positions = rng.choice(
        np.array(["Common Equity", "Preferred Equity", "Senior Debt", "Mezzanine Debt", "Fund Interest"], dtype=object),
        size=record_count,
        p=[0.42, 0.18, 0.2, 0.12, 0.08],
    )
    security_types = np.where(np.isin(positions, ["Senior Debt", "Mezzanine Debt"]), "Debt", "Equity")
    initial_cost = np.round(rng.uniform(250_000, 25_000_000, size=record_count), 2)
    ownership = np.round(rng.uniform(0.005, 0.35, size=record_count), 6)
    prices = rng.uniform(5, 250, size=record_count)
    shares = np.round(initial_cost / prices, 4)
    reporting_day = np.datetime64(reporting_date, "D")
    investment_dates = (reporting_day - rng.integers(365, 365 * 12, size=record_count).astype("timedelta64[D]")).astype(str)
    acquisition_dates = (investment_dates.astype("datetime64[D]") - rng.integers(0, 90, size=record_count).astype("timedelta64[D]")).astype(str)
    periods = _period_dates(reporting_date, period_count)

    running_cost = np.zeros(record_count, dtype=float)
    running_fv = np.zeros(record_count, dtype=float)
    running_accrued_interest = np.zeros(record_count, dtype=float)
    rf_chunks: list[pd.DataFrame] = []
    tx_chunks: list[pd.DataFrame] = []
    reference_number = 1

    for period_index, period_end in enumerate(periods):
        beginning_cost = running_cost.copy()
        beginning_fv = running_fv.copy()
        beginning_accrued_interest = running_accrued_interest.copy()
        count = record_count
        kind = np.empty(count, dtype=object)
        purchase = np.zeros(count)
        follow_on = np.zeros(count)
        capitalized_interest = np.zeros(count)
        disposals = np.zeros(count)
        writeoffs = np.zeros(count)
        adjustment = np.zeros(count)
        other_adjustment = np.zeros(count)
        unrealized = np.zeros(count)
        interest_income = np.zeros(count)
        ending_accrued_interest = beginning_accrued_interest.copy()
        realized = np.zeros(count)
        fx = np.zeros(count)
        write_up = np.zeros(count)
        write_down = np.zeros(count)
        carrying_value = np.zeros(count)
        transaction_amount = np.zeros(count)
        fair_value_change = np.zeros(count)

        if period_index == 0:
            kind[:] = "Investment Purchase"
            purchase[:] = initial_cost
            transaction_amount[:] = initial_cost
            fair_value_change[:] = initial_cost
        else:
            event_types = np.array([
                "Investment Purchase", "Follow-on Investment", "Investment Disposal", "Write-off",
                "Transfer/Adjustment", "Unrealized Gain", "FX Gain/Loss", "Write-up", "Write-down",
                "PIK Interest", "Interest Income", "Dividend Income", "Accretion/Amortization",
                "Maturity/Redemption", "Principal Repayment",
            ], dtype=object)
            profile_weights = {
                "PUBLIC_EQUITY": [0.1, 0.08, 0.12, 0.01, 0.03, 0.28, 0.04, 0.08, 0.04, 0.0, 0.0, 0.22, 0.0, 0.0, 0.0],
                "FIXED_INCOME": [0.05, 0.04, 0.1, 0.02, 0.04, 0.16, 0.05, 0.05, 0.04, 0.08, 0.25, 0.02, 0.08, 0.025, 0.045],
                "PRIVATE_CREDIT": [0.06, 0.1, 0.12, 0.05, 0.03, 0.15, 0.04, 0.03, 0.02, 0.16, 0.16, 0.0, 0.01, 0.06, 0.06],
                "ALTERNATIVES": [0.1, 0.1, 0.08, 0.03, 0.05, 0.28, 0.06, 0.1, 0.08, 0.02, 0.04, 0.06, 0.01, 0.0, 0.0],
            }
            weights = np.asarray(profile_weights.get(profile.upper(), [0.1, 0.16, 0.1, 0.025, 0.045, 0.36, 0.08, 0.07, 0.06, 0.02, 0.05, 0.01, 0.005, 0.005, 0.005]), dtype=float)
            weights /= weights.sum()
            event_indexes = rng.choice(len(event_types), size=count, p=weights)
            kind[:] = event_types[event_indexes]
            no_cost = beginning_cost <= 0
            reentry = no_cost & ~np.isin(kind, ["Investment Purchase", "Follow-on Investment"])
            kind[reentry] = "Investment Purchase"
            no_fair_value_events = (beginning_fv <= 0) & np.isin(kind, ["Unrealized Gain", "FX Gain/Loss", "Write-up", "Write-down", "Dividend Income"])
            kind[no_fair_value_events] = "Write-off"
            purchase_mask = kind == "Investment Purchase"
            follow_mask = kind == "Follow-on Investment"
            disposal_mask = kind == "Investment Disposal"
            writeoff_mask = kind == "Write-off"
            adjust_mask = kind == "Transfer/Adjustment"
            unrealized_mask = kind == "Unrealized Gain"
            fx_mask = kind == "FX Gain/Loss"
            writeup_mask = kind == "Write-up"
            writedown_mask = kind == "Write-down"
            pik_mask = kind == "PIK Interest"
            interest_mask = kind == "Interest Income"
            dividend_mask = kind == "Dividend Income"
            accretion_mask = kind == "Accretion/Amortization"
            maturity_mask = kind == "Maturity/Redemption"
            repayment_mask = kind == "Principal Repayment"

            purchase[purchase_mask] = np.round(np.maximum(1000, beginning_cost[purchase_mask] * rng.uniform(0.03, 0.16, purchase_mask.sum())), 2)
            follow_on[follow_mask] = np.round(np.maximum(1000, beginning_cost[follow_mask] * rng.uniform(0.03, 0.2, follow_mask.sum())), 2)
            sale_mask = disposal_mask | maturity_mask | repayment_mask
            disposal_fraction = rng.uniform(0.05, 0.3, sale_mask.sum())
            disposal_fraction[maturity_mask[sale_mask]] = 1.0
            disposal_fraction[repayment_mask[sale_mask]] = rng.uniform(0.1, 0.5, repayment_mask.sum())
            disposals[sale_mask] = np.round(beginning_cost[sale_mask] * disposal_fraction, 2)
            carrying_value[sale_mask] = np.round(beginning_fv[sale_mask] * disposal_fraction, 2)
            realized[sale_mask] = np.round(carrying_value[sale_mask] * rng.uniform(-0.2, 0.25, sale_mask.sum()), 2)
            writeoff_fraction = rng.uniform(0.02, 0.1, writeoff_mask.sum())
            writeoffs[writeoff_mask] = np.round(beginning_cost[writeoff_mask] * writeoff_fraction, 2)
            carrying_value[writeoff_mask] = np.round(beginning_fv[writeoff_mask] * writeoff_fraction, 2)
            adjustment[adjust_mask] = np.round(beginning_cost[adjust_mask] * rng.uniform(-0.04, 0.04, adjust_mask.sum()), 2)
            adjustment[adjust_mask] = np.maximum(adjustment[adjust_mask], -beginning_cost[adjust_mask] * 0.2)
            unrealized[unrealized_mask] = np.round(beginning_fv[unrealized_mask] * rng.uniform(-0.1, 0.12, unrealized_mask.sum()), 2)
            fx[fx_mask] = np.round(beginning_fv[fx_mask] * rng.uniform(-0.035, 0.035, fx_mask.sum()), 2)
            write_up[writeup_mask] = np.round(beginning_fv[writeup_mask] * rng.uniform(0.005, 0.05, writeup_mask.sum()), 2)
            write_down[writedown_mask] = np.round(beginning_fv[writedown_mask] * rng.uniform(0.005, 0.05, writedown_mask.sum()), 2)
            capitalized_interest[pik_mask] = np.round(beginning_cost[pik_mask] * rng.uniform(0.002, 0.015, pik_mask.sum()), 2)
            interest_income[interest_mask] = np.round(beginning_cost[interest_mask] * rng.uniform(0.005, 0.025, interest_mask.sum()), 2)
            ending_accrued_interest = np.round(beginning_accrued_interest + interest_income, 2)
            other_adjustment[accretion_mask] = np.round(beginning_cost[accretion_mask] * rng.uniform(-0.015, 0.02, accretion_mask.sum()), 2)
            transaction_amount[dividend_mask] = np.round(beginning_fv[dividend_mask] * rng.uniform(0.002, 0.02, dividend_mask.sum()), 2)
            transaction_amount[purchase_mask] = purchase[purchase_mask]
            transaction_amount[follow_mask] = follow_on[follow_mask]
            transaction_amount[sale_mask] = np.maximum(0, carrying_value[sale_mask] + realized[sale_mask])
            transaction_amount[writeoff_mask] = writeoffs[writeoff_mask]
            transaction_amount[adjust_mask] = np.abs(adjustment[adjust_mask])
            transaction_amount[unrealized_mask] = np.abs(unrealized[unrealized_mask])
            transaction_amount[fx_mask] = np.abs(fx[fx_mask])
            transaction_amount[writeup_mask] = write_up[writeup_mask]
            transaction_amount[writedown_mask] = write_down[writedown_mask]
            transaction_amount[pik_mask] = capitalized_interest[pik_mask]
            transaction_amount[interest_mask] = interest_income[interest_mask]
            transaction_amount[accretion_mask] = np.abs(other_adjustment[accretion_mask])
            fair_value_change = purchase + follow_on + capitalized_interest + interest_income - carrying_value + adjustment + other_adjustment + unrealized + fx + write_up - write_down

        ending_cost = np.round(beginning_cost + purchase + follow_on + capitalized_interest - disposals - writeoffs + adjustment + other_adjustment, 2)
        ending_fv = np.round(beginning_fv + fair_value_change, 2)
        running_cost = np.maximum(ending_cost, 0)
        running_fv = np.maximum(ending_fv, 0)
        running_accrued_interest = ending_accrued_interest
        source_refs = _numbered("PE-", reference_number, count, 10)
        batch_numbers = _numbered("BATCH-", reference_number, count, 10)
        reference_number += count
        common = {
            "Fund Name": fund_names[fund_indexes], "Fund ID": fund_ids[fund_indexes],
            "Deal Name": deal_names[deal_indexes], "Investment ID": investment_ids,
            "Position": positions, "Investor Name": ref_investors[investor_indexes],
        }
        rf_chunks.append(pd.DataFrame({
            "S.No": np.arange(period_index * count + 1, (period_index + 1) * count + 1),
            **common, "Reporting Period": period_end.isoformat(), "Beginning Cost": beginning_cost,
            "Purchases/Additions": purchase, "Follow-on Investment": follow_on,
            "PIK/Capitalized Interest": capitalized_interest, "Disposals": disposals,
            "Write-offs": writeoffs, "Transfers/Adjustments": adjustment,
            "Other Adjustments": other_adjustment, "Ending Cost": ending_cost,
            "Beginning Fair Value": beginning_fv, "Interest Income": interest_income,
            "Beginning Accrued Interest": beginning_accrued_interest,
            "Ending Accrued Interest": ending_accrued_interest,
            "Unrealized Gain/Loss": unrealized, "Realized Gain/Loss": realized,
            "FX Gain/Loss": fx, "Write-up": write_up, "Write-down": -write_down,
            "Ending Fair Value": ending_fv, "Source Transaction Reference": source_refs,
        }, columns=ROLLFORWARD_COLUMNS))
        tx_chunks.append(_base_transactions(
            source_refs, batch_numbers, kind, common["Fund Name"], common["Fund ID"],
            common["Deal Name"], common["Investment ID"], common["Position"],
            common["Investor Name"], np.full(count, period_end.isoformat(), dtype=object),
            transaction_amount, purchase + follow_on + capitalized_interest - disposals - writeoffs + adjustment + other_adjustment,
            fair_value_change, realized, carrying_value,
            ref_affiliates[affiliate_indexes], ref_vendors[vendor_indexes],
            np.full(count, "BA-000000000001", dtype=object),
            np.full(count, "LOC-0001", dtype=object),
        ))

    # Fund-level activity is part of the same transaction model and uses the same references.
    fund_tx_chunks = []
    fund_kinds = ["Management Fee", "Fund Expense", "Interest Income", "Dividend Income", "Capital Contribution", "Investor Distribution", "Other Fund Transaction"]
    fund_event_count = fund_count * period_count * len(fund_kinds)
    fund_numbers = np.arange(fund_event_count)
    fund_indexes_for_tx = np.repeat(np.arange(fund_count), period_count * len(fund_kinds))
    period_indexes_for_tx = np.tile(np.repeat(np.arange(period_count), len(fund_kinds)), fund_count)
    kinds_for_fund = np.tile(np.array(fund_kinds, dtype=object), fund_count * period_count)
    fund_amounts = np.round(rng.uniform(5_000, 2_000_000, size=fund_event_count), 2)
    fund_refs = _numbered("PE-", reference_number, fund_event_count, 10)
    fund_batches = _numbered("BATCH-", reference_number, fund_event_count, 10)
    fund_periods = np.array([periods[int(index)].isoformat() for index in period_indexes_for_tx], dtype=object)
    fund_investors = ref_investors[rng.integers(0, len(ref_investors), size=fund_event_count)]
    fund_affiliates = ref_affiliates[fund_indexes_for_tx % len(ref_affiliates)]
    fund_vendors = ref_vendors[rng.integers(0, len(ref_vendors), size=fund_event_count)]
    fund_tx_chunks.append(_base_transactions(
        fund_refs, fund_batches, kinds_for_fund, fund_names[fund_indexes_for_tx],
        fund_ids[fund_indexes_for_tx], np.full(fund_event_count, "", dtype=object),
        np.full(fund_event_count, "", dtype=object), np.full(fund_event_count, "", dtype=object),
        fund_investors, fund_periods, fund_amounts, np.zeros(fund_event_count),
        np.zeros(fund_event_count), np.zeros(fund_event_count), fund_amounts,
        fund_affiliates, fund_vendors, np.full(fund_event_count, "BA-000000000001", dtype=object),
        np.full(fund_event_count, "LOC-0001", dtype=object),
    ))
    transactions = pd.concat([*tx_chunks, *fund_tx_chunks], ignore_index=True)
    rollforward = pd.concat(rf_chunks, ignore_index=True).sort_values(["Investment ID", "Reporting Period"], kind="stable", ignore_index=True)
    rollforward["S.No"] = np.arange(1, len(rollforward) + 1)

    latest = rollforward[rollforward["Reporting Period"] == reporting_date.isoformat()].copy()
    investment_index = latest["Investment ID"].str.removeprefix("INV-").astype(int).to_numpy() - 1
    soi = pd.DataFrame({
        "S.No": np.arange(1, len(latest) + 1), "Fund Name": latest["Fund Name"],
        "Fund ID": latest["Fund ID"], "Deal Name": latest["Deal Name"],
        "Investment ID": latest["Investment ID"], "Position": latest["Position"],
        "Security Type": security_types[investment_index], "Investor Name": latest["Investor Name"],
        "Investment Date": investment_dates[investment_index],
        "Acquisition Date": acquisition_dates[investment_index], "Currency": base_currency,
        "Beginning Cost": latest["Beginning Cost"],
        "Additions": latest["Purchases/Additions"] + latest["Follow-on Investment"] + latest["PIK/Capitalized Interest"],
        "Disposals": latest["Disposals"], "Ending Cost": latest["Ending Cost"],
        "Beginning Fair Value": latest["Beginning Fair Value"],
        "Unrealized Gain/Loss": latest["Unrealized Gain/Loss"],
        "Realized Gain/Loss": latest["Realized Gain/Loss"], "FX Gain/Loss": latest["FX Gain/Loss"],
        "Ending Fair Value": latest["Ending Fair Value"], "Ownership %": ownership[investment_index],
        "Shares/Units": shares[investment_index],
        "Valuation Method": rng.choice(["Market Approach", "Income Approach", "Recent Transaction", "NAV"], size=len(latest)),
        "Investment Status": np.where(latest["Ending Fair Value"].to_numpy() > 0, "Active", "Realized"),
        "Reporting Date": reporting_date.isoformat(),
    }, columns=SOI_COLUMNS)
    soi = soi.sort_values(["Fund Name", "Deal Name", "Investment ID"], kind="stable", ignore_index=True)
    soi["S.No"] = np.arange(1, len(soi) + 1)
    gl = _transactions_to_gl(transactions)
    fund_summary = soi.groupby(["Fund Name", "Fund ID", "Currency"], as_index=False).agg(
        Investments=("Investment ID", "count"), Beginning_Cost=("Beginning Cost", "sum"),
        Additions=("Additions", "sum"), Disposals=("Disposals", "sum"),
        Ending_Cost=("Ending Cost", "sum"), Ending_Fair_Value=("Ending Fair Value", "sum"),
    )
    fund_summary.columns = [column.replace("_", " ") for column in fund_summary.columns]
    validation = validate_linked_data(soi, rollforward, transactions, gl, fund_names, deal_names, ref_investors, positions, ref_affiliates)
    financial_model = {
        "Fund": pd.DataFrame({"Fund ID": fund_ids, "Fund Name": fund_names}),
        "Investor": pd.DataFrame({"Investor Name": pd.unique(ref_investors)}),
        "Portfolio Company": pd.DataFrame({"Portfolio Company": deal_names}),
        "Deal": pd.DataFrame({"Deal Name": deal_names, "Fund ID": fund_ids[deal_funds], "Fund Name": fund_names[deal_funds]}),
        "Investment": soi[["Investment ID", "Fund ID", "Fund Name", "Deal Name", "Position", "Investor Name"]],
        "Security": soi[["Investment ID", "Position", "Security Type", "Shares/Units", "Currency"]],
        "Position": pd.DataFrame({"Position": pd.unique(positions)}),
        "Reporting Period": pd.DataFrame({"Reporting Period": [period.isoformat() for period in periods]}),
        "Valuation": soi[["Investment ID", "Reporting Date", "Ending Cost", "Ending Fair Value", "Valuation Method"]],
        "Transaction": transactions,
        "Accounting Event": transactions,
    }
    return {
        "soi": soi, "rollforward": rollforward, "gl": gl, "transactions": transactions,
        "fund_summary": fund_summary, "validation": validation,
        "financial_model": financial_model, "canonical_soi": soi, "profile": profile.upper(),
        "reporting_date": reporting_date.isoformat(), "base_currency": base_currency,
    }


def _transactions_to_gl(transactions: pd.DataFrame) -> pd.DataFrame:
    frames: list[pd.DataFrame] = []

    def add_line(rows: pd.DataFrame, amounts: np.ndarray, side: str, vendors: np.ndarray) -> None:
        mask = amounts > 0
        if not mask.any():
            return
        selected = rows.loc[mask]
        values = amounts[mask]
        frame = selected[["Reference Number", "Batch Number", "Affiliate", "Deal Name", "Position", "Bank Account", "Investor Name", "Fund Name"]].copy()
        frame.insert(0, "S.No", 0)
        frame["Vendor"] = vendors[mask]
        frame["Amount"] = values
        frame["Debit"] = values if side == "Debit" else 0.0
        frame["Credit"] = values if side == "Credit" else 0.0
        frame["Location ID"] = selected["Location ID"].to_numpy()
        frames.append(frame[GL_COLUMNS])

    def add_journal(rows: pd.DataFrame, debit: np.ndarray, credit: np.ndarray, debit_vendor: np.ndarray, credit_vendor: np.ndarray) -> None:
        add_line(rows, debit, "Debit", debit_vendor)
        add_line(rows, credit, "Credit", credit_vendor)

    types = transactions["Transaction Type"].to_numpy()
    amounts = transactions["Transaction Amount"].to_numpy(dtype=float)
    cost = transactions["Cost Change"].to_numpy(dtype=float)
    fv = transactions["Fair Value Change"].to_numpy(dtype=float)
    carry = transactions["Carrying Value"].to_numpy(dtype=float)
    realized = transactions["Realized Gain/Loss"].to_numpy(dtype=float)

    for kind in ["Investment Purchase", "Follow-on Investment"]:
        mask = types == kind
        rows = transactions.loc[mask]
        value = amounts[mask]
        add_journal(rows, value, value, rows["Deal Name"].to_numpy(), np.full(len(rows), "Fund Cash", dtype=object))
    for kind, credit_account in [("PIK Interest", "PIK Interest Income"), ("Interest Income", "Interest Income"), ("Dividend Income", "Dividend Income")]:
        mask = types == kind
        if mask.any():
            rows = transactions.loc[mask]
            value = amounts[mask]
            add_journal(rows, value, value, np.full(len(rows), "Investment/Accrued Interest", dtype=object), np.full(len(rows), credit_account, dtype=object))
    disposal = np.isin(types, ["Investment Disposal", "Maturity/Redemption", "Principal Repayment"])
    rows = transactions.loc[disposal]
    if disposal.any():
        proceeds = amounts[disposal]
        carrying = carry[disposal]
        gain = realized[disposal]
        add_journal(rows, proceeds, carrying, np.full(len(rows), "Fund Cash", dtype=object), rows["Deal Name"].to_numpy())
        add_line(rows, np.maximum(-gain, 0), "Debit", np.full(len(rows), "Realized Loss", dtype=object))
        add_line(rows, np.maximum(gain, 0), "Credit", np.full(len(rows), "Realized Gain", dtype=object))
    for kind in ["Write-off", "Transfer/Adjustment", "Accretion/Amortization", "Unrealized Gain", "FX Gain/Loss", "Write-up", "Write-down"]:
        mask = types == kind
        if not mask.any():
            continue
        rows = transactions.loc[mask]
        value = amounts[mask]
        signed = cost[mask] if kind in {"Transfer/Adjustment", "Accretion/Amortization"} else fv[mask]
        positive = signed >= 0
        debit = np.where(positive, value, 0.0)
        credit = np.where(positive, 0.0, value)
        add_journal(rows, debit, credit, rows["Deal Name"].to_numpy(), np.full(len(rows), kind, dtype=object))
        add_journal(rows, credit, debit, np.full(len(rows), kind, dtype=object), rows["Deal Name"].to_numpy())

    fund_patterns = {
        "Management Fee": ("Fund Expense", "Fund Cash", "debit"),
        "Fund Expense": ("Fund Expense", "Fund Cash", "debit"),
        "Interest Income": ("Fund Cash", "Interest Income", "debit"),
        "Dividend Income": ("Fund Cash", "Dividend Income", "debit"),
        "Capital Contribution": ("Fund Cash", "Partner Capital", "debit"),
        "Investor Distribution": ("Investor Distribution", "Fund Cash", "debit"),
        "Other Fund Transaction": ("Other Fund Expense", "Fund Cash", "debit"),
    }
    for kind, (debit_account, credit_account, _) in fund_patterns.items():
        mask = types == kind
        if not mask.any():
            continue
        rows = transactions.loc[mask]
        value = amounts[mask]
        add_journal(rows, value, value, np.full(len(rows), debit_account, dtype=object), np.full(len(rows), credit_account, dtype=object))

    gl = pd.concat(frames, ignore_index=True)
    gl["S.No"] = np.arange(1, len(gl) + 1)
    gl["Amount"] = gl["Amount"].round(2)
    gl["Debit"] = gl["Debit"].round(2)
    gl["Credit"] = gl["Credit"].round(2)
    return gl[GL_COLUMNS]


def validate_linked_data(
    soi: pd.DataFrame,
    rollforward: pd.DataFrame,
    transactions: pd.DataFrame,
    gl: pd.DataFrame,
    funds: np.ndarray,
    deals: np.ndarray,
    investors: np.ndarray,
    positions: np.ndarray,
    affiliates: np.ndarray,
) -> pd.DataFrame:
    latest = rollforward.sort_values("Reporting Period").groupby("Investment ID", sort=False).tail(1).set_index("Investment ID")
    soi_index = soi.set_index("Investment ID")
    cost_diff = (soi_index["Ending Cost"] - latest.loc[soi_index.index, "Ending Cost"]).abs()
    fv_diff = (soi_index["Ending Fair Value"] - latest.loc[soi_index.index, "Ending Fair Value"]).abs()
    by_batch = gl.groupby("Batch Number", sort=False).agg(Debit=("Debit", "sum"), Credit=("Credit", "sum"))
    refs_in_gl = set(gl["Reference Number"].unique())
    all_refs_present = transactions["Reference Number"].isin(refs_in_gl).all()
    ordered = rollforward.sort_values(["Investment ID", "Reporting Period"], kind="stable")
    prev_cost = ordered.groupby("Investment ID", sort=False)["Ending Cost"].shift()
    prev_fv = ordered.groupby("Investment ID", sort=False)["Ending Fair Value"].shift()
    continuity = ((prev_cost.isna()) | np.isclose(ordered["Beginning Cost"], prev_cost, atol=0.01)) & ((prev_fv.isna()) | np.isclose(ordered["Beginning Fair Value"], prev_fv, atol=0.01))
    expected_cost = (
        ordered["Beginning Cost"] + ordered["Purchases/Additions"] + ordered["Follow-on Investment"]
        + ordered["PIK/Capitalized Interest"] - ordered["Disposals"] - ordered["Write-offs"]
        + ordered["Transfers/Adjustments"] + ordered["Other Adjustments"]
    )
    cost_formula = np.isclose(expected_cost, ordered["Ending Cost"], atol=0.01).all()
    transaction_fv_changes = transactions.set_index("Reference Number")["Fair Value Change"]
    expected_fv = ordered["Beginning Fair Value"] + ordered["Source Transaction Reference"].map(transaction_fv_changes)
    fv_formula = np.isclose(expected_fv, ordered["Ending Fair Value"], atol=0.01).all()
    valid_dates = pd.to_datetime(transactions["Reporting Period"], errors="coerce").notna().all()
    ref_batch_consistency = transactions.groupby("Reference Number", sort=False)["Batch Number"].nunique().le(1).all()
    gl_ref_batch_consistency = gl.groupby("Reference Number", sort=False)["Batch Number"].nunique().le(1).all()
    valid_relationships = soi[["Fund Name", "Deal Name"]].drop_duplicates().merge(
        rollforward[["Fund Name", "Deal Name"]].drop_duplicates(),
        on=["Fund Name", "Deal Name"], how="left", indicator=True,
    )["_merge"].eq("both").all()
    checks = [
        ("SOI ending cost reconciles to RollForward", bool(cost_diff.le(0.01).all()), f"Maximum difference: {cost_diff.max():,.2f}"),
        ("SOI ending fair value reconciles to RollForward", bool(fv_diff.le(0.01).all()), f"Maximum difference: {fv_diff.max():,.2f}"),
        ("RollForward transactions reconcile to GL", bool(all_refs_present and rollforward["Source Transaction Reference"].isin(refs_in_gl).all()), f"{transactions['Reference Number'].nunique():,} source transactions"),
        ("GL debits equal credits by batch", bool(np.isclose(by_batch["Debit"], by_batch["Credit"], atol=0.01).all()), f"{len(by_batch):,} batches"),
        ("GL debits equal credits overall", bool(np.isclose(gl["Debit"].sum(), gl["Credit"].sum(), atol=0.01)), f"Debit {gl['Debit'].sum():,.2f}; Credit {gl['Credit'].sum():,.2f}"),
        ("Beginning balance equals previous period ending", bool(continuity.all()), f"{len(ordered):,} rollforward rows"),
        ("RollForward ending cost formula", bool(cost_formula), "Beginning + purchases + follow-ons - disposals - write-offs +/- adjustments"),
        ("RollForward ending fair-value formula", bool(fv_formula), "Beginning fair value + source transaction fair-value movement"),
        ("Accrued interest carries forward", bool(((ordered.groupby("Investment ID", sort=False)["Ending Accrued Interest"].shift().isna()) | np.isclose(ordered["Beginning Accrued Interest"], ordered.groupby("Investment ID", sort=False)["Ending Accrued Interest"].shift(), atol=0.01)).all()), "Prior ending accrued interest equals current beginning accrued interest"),
        ("Accrued interest formula", bool(np.isclose(ordered["Beginning Accrued Interest"] + ordered["Interest Income"], ordered["Ending Accrued Interest"], atol=0.01).all()), "Beginning accrued interest + interest income"),
        ("Fund names are valid", bool(soi["Fund Name"].isin(funds).all() and rollforward["Fund Name"].isin(funds).all() and gl["Fund Name"].isin(funds).all()), f"{len(funds):,} funds"),
        ("Deal names are valid", bool(soi["Deal Name"].isin(deals).all() and rollforward["Deal Name"].isin(deals).all() and gl.loc[gl["Deal Name"].ne(""), "Deal Name"].isin(deals).all()), f"{len(deals):,} deals; fund-level GL rows may be blank"),
        ("Investor names are valid", bool(soi["Investor Name"].isin(investors).all() and rollforward["Investor Name"].isin(investors).all() and gl["Investor Name"].isin(investors).all()), "Investor master values"),
        ("Position values are valid", bool(soi["Position"].isin(positions).all() and rollforward["Position"].isin(positions).all() and gl.loc[gl["Position"].ne(""), "Position"].isin(positions).all()), "Fund-level GL rows may be blank"),
        ("Affiliate names are valid", bool(gl["Affiliate"].isin(affiliates).all()), f"{len(affiliates):,} affiliates"),
        ("Transaction reference numbers are unique", bool(transactions["Reference Number"].is_unique), f"{len(transactions):,} transactions"),
        ("Batch numbers are unique", bool(transactions["Batch Number"].is_unique), f"{transactions['Batch Number'].nunique():,} batches"),
        ("Accounting dates are valid", bool(valid_dates), "All source transaction periods parse as dates"),
        ("Reference numbers map to one batch", bool(ref_batch_consistency and gl_ref_batch_consistency), "Journal lines may repeat references only within one batch"),
        ("Fund and deal relationships are valid", bool(valid_relationships), "Fund/deal pairs agree between SOI and RollForward"),
    ]
    return pd.DataFrame(checks, columns=["Validation", "Status", "Details"])


def export_soi_excel(data: dict, output_path: str | Path | None = None) -> bytes | Path:
    workbook = Workbook(write_only=True)
    header_fill = PatternFill("solid", fgColor="17365D")
    header_font = Font(color="FFFFFF", bold=True)
    section_fill = PatternFill("solid", fgColor="DCE6F1")
    section_font = Font(bold=True, color="17365D")
    columns = data.get("soi_columns", SOI_COLUMNS)
    hierarchy = [column for column in data.get("hierarchy", ["Fund Name"]) if column in columns]
    profile_settings = data.get("profile_settings", {})
    title = profile_settings.get("title", "Schedule of Investments")
    cost_fields = set(data.get("cost_fields", ["Beginning Cost", "Additions", "Disposals", "Ending Cost"]))
    fair_value_fields = set(data.get("fair_value_fields", ["Beginning Fair Value", "Ending Fair Value"]))
    currency_fields = cost_fields | fair_value_fields | set(data.get("currency_fields", []))
    percentage_fields = set(data.get("percentage_fields", ["Ownership %"]))
    date_fields = set(data.get("date_fields", ["Investment Date", "Acquisition Date", "Reporting Date"]))
    total_fields = [column for column in columns if column in currency_fields]
    scale = float(data.get("display_scale", 1))
    scale_suffix = "," * (1 if scale == 1000 else 2 if scale == 1_000_000 else 0)
    currency = data.get("base_currency", "USD")
    detail = workbook.create_sheet("SOI Detail")
    detail.freeze_panes = "A5"
    detail.append([title])
    detail.append([f"Reporting Date: {data['reporting_date']} | Base Currency: {currency} | Display scale: {scale:g}"])
    canonical = data.get("canonical_soi", data["soi"])
    detail.append([f"Funds: {canonical['Fund Name'].nunique():,} | Investments: {len(data['soi']):,}"])
    header_cells = []
    for value in columns:
        cell = WriteOnlyCell(detail, value=value)
        cell.fill = header_fill
        cell.font = header_font
        header_cells.append(cell)
    detail.append(header_cells)
    detail.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{len(data['soi']) + 4}"
    for index, column in enumerate(columns, start=1):
        detail.column_dimensions[get_column_letter(index)].width = min(max(len(column) + 2, 13), 25)
    rows_to_render = data["soi"]
    group_fields = hierarchy[:2]
    grouped = rows_to_render.groupby(group_fields, sort=False, dropna=False) if group_fields else [("", rows_to_render)]
    row_number = 4
    for group_key, group in grouped:
        key_values = group_key if isinstance(group_key, tuple) else (group_key,)
        label = " | ".join(f"{name}: {value}" for name, value in zip(group_fields, key_values))
        section = [None] * len(columns)
        section[0] = label
        section_cells = []
        for value in section:
            cell = WriteOnlyCell(detail, value=value)
            cell.fill = section_fill
            cell.font = section_font
            section_cells.append(cell)
        detail.append(section_cells)
        row_number += 1
        for row in group[columns].itertuples(index=False, name=None):
            output = []
            for column, value in zip(columns, row):
                if column in date_fields and not pd.isna(value):
                    value = pd.Timestamp(value).to_pydatetime()
                cell = WriteOnlyCell(detail, value=None if pd.isna(value) else value)
                if column in currency_fields:
                    cell.number_format = f'"{currency}" #,##0.00{scale_suffix};[Red]("{currency}" #,##0.00{scale_suffix});-'
                elif column in percentage_fields:
                    cell.number_format = '0.00%'
                elif column in date_fields:
                    cell.number_format = 'yyyy-mm-dd'
                elif column in {"Shares", "Shares/Units", "Par Amount"}:
                    cell.number_format = '#,##0.0000;[Red](#,##0.0000);-'
                elif column == "Interest Rate":
                    cell.number_format = '0.00%'
                output.append(cell)
            detail.append(output)
            row_number += 1
        subtotal = [None] * len(columns)
        subtotal[0] = f"Subtotal: {label}"
        for index, column in enumerate(columns):
            if column in total_fields:
                subtotal[index] = group[column].sum()
        subtotal_cells = []
        for index, value in enumerate(subtotal):
            cell = WriteOnlyCell(detail, value=value)
            cell.font = Font(bold=True)
            if columns[index] in currency_fields and value is not None:
                cell.number_format = f'"{currency}" #,##0.00{scale_suffix};[Red]("{currency}" #,##0.00{scale_suffix});-'
            subtotal_cells.append(cell)
        detail.append(subtotal_cells)
        row_number += 1
    grand_total = [None] * len(columns)
    grand_total[0] = "Grand Total"
    for index, column in enumerate(columns):
        if column in total_fields:
            grand_total[index] = rows_to_render[column].sum()
    detail.append(grand_total)
    row_number += 1
    detail.auto_filter.ref = f"A4:{get_column_letter(len(columns))}{len(data['soi']) + 4}"

    summary = workbook.create_sheet("Fund Summary")
    summary.freeze_panes = "A2"
    summary_headers = []
    for value in data["fund_summary"].columns:
        cell = WriteOnlyCell(summary, value=value)
        cell.fill = header_fill
        cell.font = header_font
        summary_headers.append(cell)
    summary.append(summary_headers)
    for row in data["fund_summary"].itertuples(index=False, name=None):
        output = []
        for column, value in zip(data["fund_summary"].columns, row):
            cell = WriteOnlyCell(summary, value=value)
            if column in {"Beginning Cost", "Additions", "Disposals", "Ending Cost", "Ending Fair Value"}:
                currency = data["base_currency"]
                cell.number_format = f'"{currency}" #,##0.00;[Red]("{currency}" #,##0.00);-'
            output.append(cell)
        summary.append(output)
    summary.auto_filter.ref = f"A1:{get_column_letter(len(data['fund_summary'].columns))}{len(data['fund_summary']) + 1}"
    validation = workbook.create_sheet("Validation")
    validation.append(list(data["validation"].columns))
    for row in data["validation"].itertuples(index=False, name=None):
        validation.append(list(row))
    footnote_columns = [column for column in data.get("footnote_fields", []) if column in columns]
    if footnote_columns:
        footnotes = workbook.create_sheet("Footnotes")
        footnotes.append(["Marker", "Description"])
        footnotes.append(["1", "Reported using a valuation practical expedient or NAV estimate."])
        footnotes.append(["2", "Investment is in a realization, restructuring, or non-accrual status."])
    output = Path(output_path) if output_path is not None else BytesIO()
    if isinstance(output, Path):
        output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    return output if isinstance(output, Path) else output.getvalue()


def export_soi_pdf(data: dict, output_path: str | Path | None = None) -> bytes | Path:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A3, landscape, legal
    from reportlab.pdfbase.pdfmetrics import stringWidth
    from reportlab.pdfgen import canvas

    output = Path(output_path) if output_path is not None else BytesIO()
    output.parent.mkdir(parents=True, exist_ok=True) if isinstance(output, Path) else None
    columns = data.get("soi_columns", SOI_COLUMNS)
    page_width, page_height = landscape(A3 if len(columns) > 12 else legal)
    pdf = canvas.Canvas(str(output) if isinstance(output, Path) else output, pagesize=(page_width, page_height), pageCompression=1)
    margin = 28
    page_number = 0
    hierarchy = [column for column in data.get("hierarchy", ["Fund Name"]) if column in columns]
    profile_settings = data.get("profile_settings", {})
    title = profile_settings.get("title", "Schedule of Investments")
    currency_fields = set(data.get("cost_fields", ["Beginning Cost", "Additions", "Disposals", "Ending Cost"])) | set(data.get("fair_value_fields", ["Beginning Fair Value", "Ending Fair Value"])) | set(data.get("currency_fields", []))
    percentage_fields = set(data.get("percentage_fields", ["Ownership %"]))
    footnote_fields = [column for column in data.get("footnote_fields", []) if column in columns]
    total_fields = [column for column in columns if column in currency_fields]
    currency = data.get("base_currency", "USD")
    scale = float(data.get("display_scale", 1))
    scale_label = "in thousands" if scale == 1000 else "in millions" if scale == 1_000_000 else ""
    weights = [1.8 if any(token in column.lower() for token in ["name", "company", "description", "security"]) else 1.25 if column in currency_fields else 0.8 for column in columns]
    available_width = page_width - margin * 2
    widths = [available_width * weight / sum(weights) for weight in weights]
    y = 0

    def draw_page_header() -> float:
        nonlocal page_number
        page_number += 1
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.rect(0, page_height - 72, page_width, 72, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 15)
        pdf.drawString(margin, page_height - 25, title)
        pdf.setFont("Helvetica", 9)
        canonical = data.get("canonical_soi", data["soi"])
        fund_names = list(canonical["Fund Name"].drop_duplicates().head(3))
        fund_label = "; ".join(fund_names)
        if canonical["Fund Name"].nunique() > len(fund_names):
            fund_label += f" (+{canonical['Fund Name'].nunique() - len(fund_names):,} more)"
        pdf.drawString(margin, page_height - 43, f"Reporting date: {data['reporting_date']}    Currency: {currency}    {scale_label}    Fund(s): {fund_label[:120]}")
        pdf.setFillColor(colors.black)
        y_value = page_height - 89
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.rect(margin, y_value - 4, sum(widths), 20, fill=1, stroke=0)
        pdf.setFillColor(colors.white)
        pdf.setFont("Helvetica-Bold", 7)
        x = margin + 3
        for column, width in zip(columns, widths):
            pdf.drawString(x, y_value + 2, column[:24])
            x += width
        pdf.setFillColor(colors.black)
        return y_value - 10

    y = draw_page_header()
    sums = {column: float(data["soi"][column].sum()) for column in total_fields}
    row_count = 0

    def new_page() -> None:
        nonlocal y
        pdf.setFont("Helvetica", 8)
        pdf.drawRightString(page_width - margin, 17, f"Page {page_number}")
        pdf.showPage()
        y = draw_page_header()

    def wrap_text(value: str, width: float, font: str, font_size: float) -> list[str]:
        lines: list[str] = []
        current = ""
        for word in str(value).split():
            candidate = f"{current} {word}".strip()
            if current and stringWidth(candidate, font, font_size) > width:
                lines.append(current)
                current = word
            else:
                current = candidate
        if current:
            lines.append(current)
        return lines or [""]

    def draw_row(values: list[Any], bold: bool = False, fill: Any = None) -> None:
        nonlocal y
        font = "Helvetica-Bold" if bold else "Helvetica"
        font_size = 7
        wrapped = [wrap_text(value, width - 6, font, font_size) for value, width in zip(values, widths)]
        row_height = max(12, max(map(len, wrapped), default=1) * 8 + 4)
        if y - row_height < margin + 22:
            new_page()
        if fill is not None:
            pdf.setFillColor(fill)
            pdf.rect(margin, y - row_height + 3, available_width, row_height, fill=1, stroke=0)
            pdf.setFillColor(colors.black)
        pdf.setFont(font, font_size)
        x = margin + 3
        for lines, width in zip(wrapped, widths):
            baseline = y - 6
            for line in lines:
                pdf.drawString(x, baseline, line)
                baseline -= 8
            x += width
        y -= row_height

    def formatted_value(column: str, value: Any) -> str:
        if pd.isna(value):
            return ""
        if column in currency_fields and isinstance(value, (int, float, np.number)):
            return f"{currency} {float(value) / scale:,.2f}"
        if column in percentage_fields and isinstance(value, (int, float, np.number)):
            return f"{float(value):.2%}"
        if column in data.get("date_fields", []) and not isinstance(value, (date, pd.Timestamp)):
            return pd.Timestamp(value).strftime("%Y-%m-%d")
        if column == "Interest Rate" and isinstance(value, (int, float, np.number)):
            return f"{float(value):.2%}"
        return str(value)

    group_fields = hierarchy[:2]
    grouped = data["soi"].groupby(group_fields, sort=False, dropna=False) if group_fields else [("", data["soi"])]
    for group_key, group in grouped:
        key_values = group_key if isinstance(group_key, tuple) else (group_key,)
        label = " | ".join(f"{name}: {value}" for name, value in zip(group_fields, key_values))
        if y < margin + 40:
            new_page()
        pdf.setFillColor(colors.HexColor("#DCE6F1"))
        pdf.rect(margin, y - 12, available_width, 16, fill=1, stroke=0)
        pdf.setFillColor(colors.HexColor("#17365D"))
        pdf.setFont("Helvetica-Bold", 8)
        pdf.drawString(margin + 3, y - 7, label[:160])
        pdf.setFillColor(colors.black)
        y -= 17
        for row in group[columns].itertuples(index=False, name=None):
            draw_row([formatted_value(column, value) for column, value in zip(columns, row)])
            row_count += 1
        subtotal = [""] * len(columns)
        subtotal[0] = "Subtotal"
        for index, column in enumerate(columns):
            if column in total_fields:
                subtotal[index] = formatted_value(column, group[column].sum())
        draw_row(subtotal, bold=True)

    grand_total = [""] * len(columns)
    grand_total[0] = f"Grand Total ({row_count:,} investments)"
    for index, column in enumerate(columns):
        if column in total_fields:
            grand_total[index] = formatted_value(column, sums[column])
    draw_row(grand_total, bold=True)
    if footnote_fields:
        markers = pd.unique(data["soi"][footnote_fields].values.ravel("K"))
        markers = [str(value) for value in markers if pd.notna(value) and str(value).strip() and str(value) not in {"No", "N/A"}]
        if markers:
            if y < 90:
                new_page()
            y -= 10
            pdf.setFont("Helvetica-Bold", 8)
            pdf.drawString(margin, y, "Footnotes")
            for marker in markers:
                y -= 11
                pdf.setFont("Helvetica", 7)
                description = "Measured using a NAV/practical expedient." if marker == "Practical Expedient" else "Reported valuation or investment-status indicator."
                draw_row([f"{marker}: {description}"] + [""] * (len(columns) - 1))
    pdf.setFont("Helvetica", 8)
    pdf.drawRightString(page_width - margin, 17, f"Page {page_number}")
    pdf.save()
    return output if isinstance(output, Path) else output.getvalue()


def dataframe_bytes(frame: pd.DataFrame, fmt: str, sheet_name: str = "Data") -> bytes:
    output = BytesIO()
    if fmt == "csv":
        frame.to_csv(output, index=False)
    elif fmt == "parquet":
        frame.to_parquet(output, index=False)
    elif fmt == "json":
        frame.to_json(output, orient="records", lines=True)
    elif fmt == "xlsx":
        with pd.ExcelWriter(output, engine="openpyxl") as writer:
            max_rows = 1_048_575
            for part, start in enumerate(range(0, len(frame), max_rows), start=1):
                frame.iloc[start:start + max_rows].to_excel(writer, index=False, sheet_name=f"{sheet_name[:25]} {part}")
    else:
        raise ValueError(f"Unsupported output format: {fmt}")
    return output if isinstance(output, Path) else output.getvalue()